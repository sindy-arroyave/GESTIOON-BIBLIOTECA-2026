from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
CAPTURES = DOCS / "capturas"
DIAGRAMS = DOCS / "diagramas"
OUTPUT = DOCS / "INFORME_APA7.docx"

NAVY = "12202F"
NAVY_LIGHT = "1C3044"
TEAL = "00A896"
PALE = "E8F7F4"
LIGHT = "F4F7FA"
GRID = "9AA7B2"
TEXT = "222F3E"
MUTED = "5F6B75"
WHITE = "FFFFFF"


def font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


F_TITLE = font(42, True)
F_HEAD = font(30, True)
F_BODY = font(24)
F_SMALL = font(20)
F_SMALL_B = font(20, True)
F_TINY = font(17)


def multiline(draw, xy, text, fnt, fill=TEXT, anchor="la", spacing=7, align="left"):
    draw.multiline_text(xy, text, font=fnt, fill="#" + fill, anchor=anchor, spacing=spacing, align=align)


def rounded_box(draw, box, title, lines, stereotype=None, header=NAVY_LIGHT, fill=WHITE, line=GRID):
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=18, fill="#" + fill, outline="#" + line, width=4)
    header_h = 100 if stereotype else 72
    draw.rounded_rectangle((x1, y1, x2, y1 + header_h), radius=18, fill="#" + header, outline="#" + header, width=2)
    draw.rectangle((x1, y1 + header_h - 18, x2, y1 + header_h), fill="#" + header)
    if stereotype:
        multiline(draw, ((x1 + x2) / 2, y1 + 24), stereotype, F_TINY, WHITE, "ma", align="center")
        multiline(draw, ((x1 + x2) / 2, y1 + 63), title, F_HEAD, WHITE, "ma", align="center")
    else:
        multiline(draw, ((x1 + x2) / 2, y1 + 38), title, F_HEAD, WHITE, "mm", align="center")
    y = y1 + header_h + 18
    for item in lines:
        multiline(draw, (x1 + 22, y), item, F_SMALL, TEXT)
        y += 35
    return box


def arrow(draw, points, color=TEAL, width=7, dashed=False):
    if dashed:
        for a, b in zip(points[:-1], points[1:]):
            x1, y1 = a
            x2, y2 = b
            length = math.hypot(x2 - x1, y2 - y1)
            if length == 0:
                continue
            ux, uy = (x2 - x1) / length, (y2 - y1) / length
            pos = 0
            while pos < length:
                end = min(pos + 18, length)
                draw.line((x1 + ux * pos, y1 + uy * pos, x1 + ux * end, y1 + uy * end), fill="#" + color, width=width)
                pos += 31
    else:
        draw.line(points, fill="#" + color, width=width, joint="curve")
    x1, y1 = points[-2]
    x2, y2 = points[-1]
    angle = math.atan2(y2 - y1, x2 - x1)
    size = 22
    left = (x2 - size * math.cos(angle - math.pi / 6), y2 - size * math.sin(angle - math.pi / 6))
    right = (x2 - size * math.cos(angle + math.pi / 6), y2 - size * math.sin(angle + math.pi / 6))
    draw.polygon([(x2, y2), left, right], fill="#" + color)


def inheritance(draw, child, parent, dashed=False):
    cx = (child[0] + child[2]) / 2
    cy = child[1]
    px = (parent[0] + parent[2]) / 2
    py = parent[3]
    mid = (cy + py) / 2
    if dashed:
        arrow(draw, [(cx, cy), (cx, mid), (px, mid), (px, py + 18)], NAVY_LIGHT, 5, True)
    else:
        draw.line([(cx, cy), (cx, mid), (px, mid), (px, py + 18)], fill="#" + NAVY_LIGHT, width=5, joint="curve")
    tri = [(px, py), (px - 18, py + 30), (px + 18, py + 30)]
    draw.polygon(tri, fill="#" + WHITE, outline="#" + NAVY_LIGHT)
    draw.line(tri + [tri[0]], fill="#" + NAVY_LIGHT, width=4)


def diamond(draw, center, toward, color=NAVY_LIGHT):
    x, y = center
    tx, ty = toward
    angle = math.atan2(ty - y, tx - x)
    along = (math.cos(angle), math.sin(angle))
    across = (-along[1], along[0])
    pts = []
    for a, b in [(0, 0), (18, 13), (36, 0), (18, -13)]:
        pts.append((x + along[0] * a + across[0] * b, y + along[1] * a + across[1] * b))
    draw.polygon(pts, fill="#" + color, outline="#" + color)


def diagram_header(draw, width, title, subtitle):
    draw.rectangle((0, 0, width, 105), fill="#" + NAVY)
    multiline(draw, (55, 28), title, F_TITLE, WHITE)
    multiline(draw, (width - 55, 53), subtitle, F_SMALL, WHITE, "rm")


def generate_architecture(path: Path):
    w, h = 2200, 1350
    img = Image.new("RGB", (w, h), "#F4F7FA")
    d = ImageDraw.Draw(img)
    diagram_header(d, w, "Diagrama UML de componentes", "Arquitectura por capas")
    d.rounded_rectangle((70, 145, 2130, 1290), radius=30, fill="#FFFFFF", outline="#9AA7B2", width=4)
    multiline(d, (105, 175), "«system» SistemaGestionBiblioteca", F_HEAD, NAVY_LIGHT)
    layers = [
        ((180, 260, 1650, 445), "«component»", "Biblioteca.Presentacion", "Windows Forms · navegación · formularios CRUD · mensajes"),
        ((180, 515, 1650, 700), "«component»", "Biblioteca.Negocio", "Servicios · validaciones · reglas de préstamo y devolución"),
        ((180, 770, 1650, 955), "«component»", "Biblioteca.Datos", "ADO.NET · repositorios · consultas parametrizadas"),
        ((180, 1025, 1650, 1210), "«database»", "BibliotecaDB", "SQL Server · tablas · vistas · restricciones · transacciones"),
    ]
    boxes = []
    for box, stereo, title, desc in layers:
        rounded_box(d, box, title, [desc], stereo, TEAL if title == "BibliotecaDB" else NAVY_LIGHT)
        boxes.append(box)
    entity = rounded_box(
        d,
        (1740, 455, 2050, 965),
        "Biblioteca.Entidades",
        ["IValidable", "EntidadBase", "Persona", "Autor", "Categoría", "Usuario", "Libro", "Préstamo", "DetallePrestamo"],
        "«component»",
        "3A627F",
    )
    for top, bottom in zip(boxes[:-1], boxes[1:]):
        arrow(d, [((top[0] + top[2]) / 2, top[3]), ((bottom[0] + bottom[2]) / 2, bottom[1])], TEAL, 8)
    for box, target_y in zip(boxes[:3], (535, 710, 885)):
        source_y = (box[1] + box[3]) / 2
        arrow(d, [(box[2], source_y), (1705, source_y), (1705, target_y), (entity[0], target_y)], "3A627F", 4, True)
    multiline(d, (1090, 1250), "Dependencia descendente entre capas; las entidades son compartidas mediante contratos del dominio.", F_SMALL, MUTED, "mm")
    img.save(path, quality=95)


def generate_class(path: Path):
    w, h = 2400, 1750
    img = Image.new("RGB", (w, h), "#F4F7FA")
    d = ImageDraw.Draw(img)
    diagram_header(d, w, "Diagrama UML de clases", "Modelo de dominio orientado a objetos")
    boxes = {}
    boxes["IValidable"] = rounded_box(d, (900, 150, 1500, 345), "IValidable", ["+ Validar(): IList<string>"], "«interface»", "3A627F")
    boxes["EntidadBase"] = rounded_box(d, (900, 445, 1500, 675), "EntidadBase", ["+ Id: int", "+ Validar(): IList<string>"], "«abstract»")
    boxes["Persona"] = rounded_box(d, (70, 820, 520, 1090), "Persona", ["+ Nombre: string", "+ Apellidos: string", "+ NombreCompleto: string", "# ValidarNombre(errores)"], "«abstract»")
    boxes["Categoria"] = rounded_box(d, (605, 820, 1055, 1055), "Categoria", ["+ Nombre: string", "+ Descripcion: string", "+ Validar(): IList<string>"], None, TEAL)
    boxes["Libro"] = rounded_box(d, (1140, 790, 1680, 1165), "Libro", ["+ Codigo: string", "+ Isbn: string", "+ Titulo: string", "+ AutorId: int", "+ CategoriaId: int", "+ CantidadTotal: int", "+ CantidadDisponible: int", "+ Disponible: bool"], None, TEAL)
    boxes["Prestamo"] = rounded_box(d, (1765, 760, 2325, 1125), "Prestamo", ["+ UsuarioId: int", "+ FechaPrestamo: DateTime", "+ FechaDevolucionEsperada: DateTime", "+ FechaDevolucionReal: DateTime?", "+ Estado: EstadoPrestamo", "+ Detalles: IList<DetallePrestamo>", "+ Prestamo()", "+ Validar(): IList<string>"], None, TEAL)
    boxes["Autor"] = rounded_box(d, (55, 1330, 470, 1630), "Autor", ["+ Codigo: string", "+ Nacionalidad: string", "+ FechaNacimiento: DateTime", "+ Validar(): IList<string>"], None, "3A627F")
    boxes["Usuario"] = rounded_box(d, (555, 1305, 1015, 1660), "Usuario", ["+ Documento: string", "+ Telefono: string", "+ Correo: string", "+ ProgramaAcademico: string", "+ Validar(): IList<string>"], None, "3A627F")
    boxes["Detalle"] = rounded_box(d, (1765, 1360, 2325, 1645), "DetallePrestamo", ["+ Id: int", "+ PrestamoId: int", "+ LibroId: int", "+ Cantidad: int"], None, "3A627F")
    inheritance(d, boxes["EntidadBase"], boxes["IValidable"], True)
    for name in ["Persona", "Categoria", "Libro", "Prestamo"]:
        inheritance(d, boxes[name], boxes["EntidadBase"])
    inheritance(d, boxes["Autor"], boxes["Persona"])
    inheritance(d, boxes["Usuario"], boxes["Persona"])
    px = (boxes["Prestamo"][0] + boxes["Prestamo"][2]) / 2
    diamond(d, (px, boxes["Prestamo"][3]), (px, boxes["Detalle"][1]))
    d.line((px + 36, boxes["Prestamo"][3], px + 36, boxes["Detalle"][1]), fill="#" + NAVY_LIGHT, width=5)
    multiline(d, (px + 48, boxes["Prestamo"][3] + 55), "1", F_SMALL_B, NAVY_LIGHT)
    multiline(d, (px + 48, boxes["Detalle"][1] - 35), "1..*", F_SMALL_B, NAVY_LIGHT)
    multiline(d, (1200, 1705), "Herencia con triángulo vacío · realización discontinua · composición con rombo sólido", F_SMALL, MUTED, "mm")
    img.save(path, quality=95)


def generate_data_model(path: Path):
    w, h = 2400, 1750
    img = Image.new("RGB", (w, h), "#F4F7FA")
    d = ImageDraw.Draw(img)
    diagram_header(d, w, "Modelo lógico de datos", "Relaciones y multiplicidades UML")
    boxes = {}
    boxes["Autores"] = rounded_box(d, (80, 165, 650, 515), "Autores", ["PK AutorId: int", "UQ Codigo: varchar(20)", "Nombre: nvarchar(80)", "Apellidos: nvarchar(100)", "Nacionalidad: nvarchar(80)", "FechaNacimiento: date"], "«table»")
    boxes["Categorias"] = rounded_box(d, (770, 165, 1340, 480), "Categorias", ["PK CategoriaId: int", "UQ Nombre: nvarchar(80)", "Descripcion: nvarchar(250)", "Activa: bit"], "«table»")
    boxes["Usuarios"] = rounded_box(d, (1510, 165, 2315, 595), "Usuarios", ["PK UsuarioId: int", "UQ Documento: varchar(25)", "Nombre: nvarchar(80)", "Apellidos: nvarchar(100)", "Telefono: varchar(20)", "UQ Correo: varchar(150)", "ProgramaAcademico: nvarchar(120)", "Activo: bit"], "«table»")
    boxes["Libros"] = rounded_box(d, (260, 720, 1030, 1190), "Libros", ["PK LibroId: int", "UQ Codigo: varchar(20)", "UQ ISBN: varchar(20)", "Titulo: nvarchar(200)", "FK AutorId: int", "FK CategoriaId: int", "CantidadTotal: int", "CantidadDisponible: int", "Activo: bit"], "«table»", TEAL)
    boxes["Prestamos"] = rounded_box(d, (1390, 735, 2240, 1160), "Prestamos", ["PK PrestamoId: int", "FK UsuarioId: int", "FechaPrestamo: datetime2", "FechaDevolucionEsperada: date", "FechaDevolucionReal: datetime2?", "Estado: varchar(12)"], "«table»", TEAL)
    boxes["Detalle"] = rounded_box(d, (525, 1350, 1260, 1660), "DetallePrestamos", ["PK DetallePrestamoId: int", "FK PrestamoId: int", "FK LibroId: int", "Cantidad: int", "UQ (PrestamoId, LibroId)"], "«table»", "3A627F")
    boxes["Devoluciones"] = rounded_box(d, (1510, 1360, 2240, 1650), "Devoluciones", ["PK DevolucionId: int", "FK/UQ PrestamoId: int", "FechaDevolucion: datetime2", "Observacion: nvarchar(300)"], "«table»", "3A627F")
    def assoc(a, b, a_mult, b_mult, points):
        d.line(points, fill="#" + NAVY_LIGHT, width=5, joint="curve")
        multiline(d, (points[0][0] + 12, points[0][1] - 28), a_mult, F_SMALL_B, NAVY_LIGHT)
        multiline(d, (points[-1][0] + 12, points[-1][1] - 28), b_mult, F_SMALL_B, NAVY_LIGHT)
    assoc("Autores", "Libros", "1", "0..*", [(365, 515), (365, 650), (470, 650), (470, 720)])
    assoc("Categorias", "Libros", "1", "0..*", [(1055, 480), (1055, 630), (820, 630), (820, 720)])
    assoc("Usuarios", "Prestamos", "1", "0..*", [(1890, 595), (1890, 735)])
    d.line([(1500, 1160), (1500, 1280), (1260, 1280), (1260, 1450)], fill="#" + NAVY_LIGHT, width=5, joint="curve")
    multiline(d, (1512, 1175), "1", F_SMALL_B, NAVY_LIGHT)
    multiline(d, (1272, 1405), "1..*", F_SMALL_B, NAVY_LIGHT)
    assoc("Libros", "Detalle", "1", "0..*", [(645, 1190), (645, 1350)])
    assoc("Prestamos", "Devoluciones", "1", "0..1", [(1820, 1160), (1820, 1360)])
    multiline(d, (1200, 1710), "Las llaves foráneas protegen el historial; no se usa eliminación en cascada.", F_SMALL, MUTED, "mm")
    img.save(path, quality=95)


def generate_sequence(path: Path):
    w, h = 2400, 1500
    img = Image.new("RGB", (w, h), "#F4F7FA")
    d = ImageDraw.Draw(img)
    diagram_header(d, w, "Diagrama UML de secuencia", "Caso de uso: registrar préstamo")
    xs = [170, 650, 1130, 1610, 2090]
    names = ["Encargado", "PrestamosForm", "PrestamoServicio", "PrestamoRepositorio", "SQL Server"]
    stereos = ["«actor»", "«boundary»", "«control»", "«repository»", "«database»"]
    for x, name, stereo in zip(xs, names, stereos):
        rounded_box(d, (x - 180, 175, x + 180, 325), name, [], stereo, NAVY_LIGHT if name != "SQL Server" else TEAL)
        y = 325
        while y < 1400:
            d.line((x, y, x, min(y + 20, 1400)), fill="#" + GRID, width=4)
            y += 36
    messages = [
        (0, 1, 410, "1. Seleccionar usuario, libro y fecha"),
        (1, 2, 515, "2. Registrar(usuarioId, libroId, fecha)"),
        (2, 2, 620, "3. Validar entidad y reglas"),
        (2, 3, 735, "4. RegistrarPrestamo(...)"),
        (3, 4, 840, "5. EXEC sp_RegistrarPrestamo"),
        (4, 4, 945, "6. [datos válidos] INSERT/UPDATE y COMMIT"),
        (4, 4, 1050, "7. [datos inválidos] ROLLBACK y THROW"),
        (4, 3, 1170, "8. PrestamoId"),
        (3, 2, 1260, "9. PrestamoId"),
        (2, 1, 1340, "10. Confirmación"),
    ]
    d.rounded_rectangle((1420, 890, 2320, 1135), radius=18, outline="#" + TEAL, width=4)
    multiline(d, (1450, 905), "«alt»", F_SMALL_B, TEAL)
    d.line((1420, 1018, 2320, 1018), fill="#" + TEAL, width=3)
    for src, dst, y, label in messages:
        x1, x2 = xs[src], xs[dst]
        if src == dst:
            if x1 > 1900:
                d.line((x1, y, x1 - 210, y, x1 - 210, y + 55, x1, y + 55), fill="#" + TEAL, width=5)
                arrow(d, [(x1 - 210, y + 55), (x1, y + 55)], TEAL, 5)
                multiline(d, (x1 - 500, y - 32), label, F_TINY, TEXT)
            else:
                d.line((x1, y, x1 + 210, y, x1 + 210, y + 55, x1, y + 55), fill="#" + TEAL, width=5)
                arrow(d, [(x1 + 210, y + 55), (x1, y + 55)], TEAL, 5)
                multiline(d, (x1 + 25, y - 32), label, F_TINY, TEXT)
        else:
            arrow(d, [(x1, y), (x2, y)], TEAL if src < dst else "3A627F", 5, src > dst)
            multiline(d, ((x1 + x2) / 2, y - 28), label, F_TINY, TEXT, "mm", align="center")
    img.save(path, quality=95)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, size=12)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    page_number(section.header.paragraphs[0])
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)
    for style_name, size, bold, italic, align in [
        ("Title", 14, True, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 1", 12, True, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 12, True, False, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 12, True, True, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = italic
        style.font.color.rgb = RGBColor.from_string(TEXT)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = align


def add_body(doc, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Inches(0.5) if indent else Inches(0)
    r = p.add_run(text)
    set_run_font(r, italic=italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def table_caption(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(f"Tabla {number}")
    set_run_font(r, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.line_spacing = 1
    p2.paragraph_format.keep_with_next = True
    p2.paragraph_format.first_line_indent = Inches(0)
    r2 = p2.add_run(title)
    set_run_font(r2, italic=True)


def add_table(doc, number, title, headers, rows, widths=None, font_size=9):
    table_caption(doc, number, title)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY_LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = 1
        r = p.add_run(value)
        set_run_font(r, size=font_size, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[i])
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if row_idx % 2 == 0:
                set_cell_shading(cell, "F8FAFB")
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
            if widths:
                cell.width = Inches(widths[i])
    note = doc.add_paragraph()
    note.paragraph_format.first_line_indent = Inches(0)
    note.paragraph_format.line_spacing = 1
    r = note.add_run("Nota. Elaboración propia a partir de los requisitos y del código fuente implementado.")
    set_run_font(r, size=10, italic=True)
    return table


def figure_caption(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Figura {number}")
    set_run_font(r, bold=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.first_line_indent = Inches(0)
    p2.paragraph_format.line_spacing = 1
    p2.paragraph_format.keep_with_next = True
    r2 = p2.add_run(title)
    set_run_font(r2, italic=True)


def add_figure(doc, number, title, image_path, note, width=6.45):
    figure_caption(doc, number, title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    pn = doc.add_paragraph()
    pn.paragraph_format.first_line_indent = Inches(0)
    pn.paragraph_format.line_spacing = 1
    pn.paragraph_format.space_after = Pt(0)
    r = pn.add_run("Nota. " + note)
    set_run_font(r, size=10, italic=True)


def add_page_break(doc):
    doc.add_page_break()


def add_reference(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r)


def build_report():
    DIAGRAMS.mkdir(parents=True, exist_ok=True)
    generate_architecture(DIAGRAMS / "uml-arquitectura.png")
    generate_class(DIAGRAMS / "uml-clases.png")
    generate_data_model(DIAGRAMS / "uml-modelo-datos.png")
    generate_sequence(DIAGRAMS / "uml-secuencia-prestamo.png")
    required = [
        CAPTURES / "01-inicio-real.jpg",
        CAPTURES / "02-libros-real.jpg",
        CAPTURES / "03-autores-real.jpg",
        CAPTURES / "04-categorias-real.jpg",
        CAPTURES / "05-usuarios-real.jpg",
        CAPTURES / "06-prestamos-real.jpg",
        CAPTURES / "07-devoluciones-real.jpg",
        CAPTURES / "08-consultas-real.jpg",
    ]
    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise FileNotFoundError("Faltan capturas reales: " + ", ".join(missing))

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Sistema de Gestión de Biblioteca"
    doc.core_properties.subject = "Informe técnico y funcional con normas APA, séptima edición"
    doc.core_properties.author = "Ronald Calvo"
    doc.core_properties.keywords = "C#, Windows Forms, SQL Server, UML, APA 7, biblioteca"

    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run("Sistema de Gestión de Biblioteca")
    set_run_font(r, size=14, bold=True)
    for line in [
        "Aplicación de escritorio con C#, Windows Forms y SQL Server",
        "Ronald Calvo",
        "Institución educativa",
        "Programación Avanzada",
        "Proyecto final",
        "1 de septiembre de 2026",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        r = p.add_run(line)
        set_run_font(r)
    add_page_break(doc)

    add_heading(doc, "Información del proyecto", 1)
    add_body(doc, "Título: Sistema de Gestión de Biblioteca", False, WD_ALIGN_PARAGRAPH.LEFT)
    add_body(doc, "Asignatura: Programación Avanzada", False, WD_ALIGN_PARAGRAPH.LEFT)
    add_body(doc, "Tipo de entrega: Proyecto final", False, WD_ALIGN_PARAGRAPH.LEFT)
    add_body(doc, "Tecnologías: C#, Windows Forms, ADO.NET, SQL Server, Git y GitHub", False, WD_ALIGN_PARAGRAPH.LEFT)
    add_body(doc, "Arquitectura: Presentación, lógica de negocio, acceso a datos y base de datos", False, WD_ALIGN_PARAGRAPH.LEFT)
    add_body(doc, "Propósito: automatizar la administración de libros, autores, categorías, usuarios, préstamos y devoluciones de una biblioteca académica.", False)
    add_body(doc, "Este documento presenta el análisis, el diseño, la implementación y la verificación de la solución. Se entrega junto con la solución de Visual Studio, los scripts SQL, el repositorio Git y las evidencias de funcionamiento obtenidas desde la aplicación real conectada a la instancia local de SQL Server.")
    add_page_break(doc)

    add_heading(doc, "Resumen", 1)
    add_body(doc, "El proyecto desarrolla una aplicación de escritorio para centralizar la gestión de una biblioteca académica. La solución implementa operaciones CRUD para libros, autores, categorías y usuarios; registra préstamos y devoluciones; ofrece búsquedas y seis consultas operativas; y protege la integridad mediante validaciones en el dominio, servicios, repositorios y SQL Server. La arquitectura separa presentación, lógica de negocio, acceso a datos y persistencia. El modelo orientado a objetos incorpora clases, objetos, constructores, propiedades, encapsulamiento, herencia, polimorfismo por medio de la interfaz IValidable y colecciones para el detalle de los préstamos. Las operaciones críticas se ejecutan mediante procedimientos almacenados transaccionales que actualizan el historial y las existencias como una unidad. La verificación comprendió compilación, pruebas automáticas, pruebas SQL y revisión visual. Las capturas incluidas fueron obtenidas de la aplicación ejecutándose contra BibliotecaDB y evidencian cuatro títulos, doce ejemplares disponibles, tres usuarios y dos préstamos activos. El resultado satisface los requerimientos funcionales y no funcionales establecidos para la actividad y queda preparado para su ejecución en Visual Studio y su publicación en GitHub.", False)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    r = p.add_run("Palabras clave: ")
    set_run_font(r, italic=True)
    r2 = p.add_run("biblioteca, C#, Windows Forms, SQL Server, arquitectura por capas, UML.")
    set_run_font(r2)
    add_page_break(doc)

    add_heading(doc, "Tabla de contenido", 1)
    contents = [
        "1. Introducción", "2. Planteamiento del problema, objetivos y alcance", "3. Análisis de requerimientos",
        "4. Casos de uso", "5. Diseño UML y arquitectura", "6. Modelo de datos y diccionario",
        "7. Implementación de módulos", "8. Validaciones y manejo de errores", "9. Evidencias reales del sistema",
        "10. Pruebas de funcionamiento", "11. Trazabilidad y control de versiones", "12. Conclusiones y recomendaciones",
        "Referencias", "Apéndice A. Ejecución en Visual Studio", "Apéndice B. Lista de entregables",
    ]
    for item in contents:
        add_body(doc, item, False, WD_ALIGN_PARAGRAPH.LEFT)
    add_page_break(doc)

    add_heading(doc, "1. Introducción", 1)
    add_body(doc, "Las bibliotecas académicas administran información que cambia constantemente: libros, autores, categorías, estudiantes, préstamos en curso y devoluciones. Cuando estos registros se llevan de forma manual, aumentan la posibilidad de duplicados, la pérdida de información, la falta de trazabilidad y las diferencias entre la existencia física y la disponibilidad reportada.")
    add_body(doc, "La solución desarrollada centraliza el proceso en una aplicación Windows Forms. La interfaz recibe y presenta la información; la capa de negocio aplica reglas; la capa de datos usa consultas parametrizadas con ADO.NET; y SQL Server conserva los registros mediante llaves, restricciones, vistas y procedimientos almacenados. Esta separación favorece claridad, mantenimiento y pruebas reproducibles (Microsoft, 2025a, 2025b).")
    add_body(doc, "La implementación prioriza la integridad, la facilidad de uso y la organización del código. Las operaciones de préstamo y devolución se ejecutan dentro de transacciones para impedir estados parciales y mantener sincronizados el historial y la disponibilidad.")

    add_heading(doc, "2. Planteamiento del problema, objetivos y alcance", 1)
    add_heading(doc, "2.1 Planteamiento del problema", 2)
    add_body(doc, "La institución educativa registra manualmente libros, usuarios, préstamos y devoluciones. Esta práctica ocasiona pérdida de información, errores, duplicidad y dificultad para consultar el estado de los préstamos. El sistema propuesto establece una fuente central y reglas verificables para cada operación.")
    add_heading(doc, "2.2 Objetivo general", 2)
    add_body(doc, "Desarrollar una aplicación de escritorio utilizando C# y SQL Server que gestione los procesos principales de una biblioteca mediante programación orientada a objetos y conexión a bases de datos.")
    add_heading(doc, "2.3 Objetivos específicos", 2)
    for item in [
        "Aplicar clases, objetos, constructores, encapsulamiento, propiedades, métodos, herencia y colecciones.",
        "Diseñar una base de datos relacional con llaves primarias, foráneas y restricciones de integridad.",
        "Implementar CRUD para libros, autores, categorías y usuarios.",
        "Registrar préstamos y devoluciones con actualización automática de existencias.",
        "Implementar búsquedas, consultas, validaciones, manejo de excepciones y arquitectura por capas.",
        "Usar Git y documentar el análisis, el diseño, la implementación y las pruebas.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2.4 Alcance", 2)
    add_body(doc, "El sistema incluye catálogo, usuarios, circulación, búsquedas, seis consultas, validaciones, interfaz gráfica, persistencia SQL Server, datos de prueba, documentación y control de versiones. No incluye autenticación por roles ni funcionalidades de pago, de acuerdo con el enunciado.")
    add_page_break(doc)

    add_heading(doc, "3. Análisis de requerimientos", 1)
    add_table(doc, 1, "Trazabilidad de requerimientos funcionales", ["Código", "Requerimiento", "Implementación", "Criterio verificado"], [
        ("RF01", "Gestión de libros", "LibrosForm · LibroServicio · LibroRepositorio", "CRUD, búsqueda y disponibilidad"),
        ("RF02", "Gestión de autores", "AutoresForm · AutorServicio · AutorRepositorio", "Código y datos personales"),
        ("RF03", "Gestión de categorías", "CategoriasForm · CategoriaServicio", "Categorías iniciales y nuevas"),
        ("RF04", "Gestión de usuarios", "UsuariosForm · UsuarioServicio", "Documento, contacto y programa"),
        ("RF05", "Préstamos", "PrestamosForm · sp_RegistrarPrestamo", "Existencia y transacción"),
        ("RF06", "Devoluciones", "DevolucionesForm · sp_RegistrarDevolucion", "Estado y disponibilidad"),
        ("RF07", "Consultas", "ConsultasForm · vistas y agregaciones", "Seis consultas operativas"),
    ], [0.55, 1.35, 2.2, 2.4], 8)
    add_heading(doc, "3.1 Requerimientos no funcionales", 2)
    for item in [
        "Interfaz amigable con navegación lateral, títulos visibles, formularios agrupados y tablas legibles.",
        "Validación de campos y prevención de duplicados mediante reglas y restricciones UNIQUE.",
        "Mensajes claros para validación, duplicidad, integridad y conexión.",
        "Código organizado con nombres descriptivos y separación por capas.",
        "Manejo de excepciones en la interfaz y reversión transaccional en SQL Server.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "3.2 Reglas de negocio", 2)
    for item in [
        "El ISBN debe contener 10 o 13 dígitos y ser único; el código y el título son obligatorios.",
        "La cantidad total debe ser mayor que cero y la disponibilidad permanecer entre cero y el total.",
        "El documento y el correo del usuario son únicos; correo y teléfono deben tener formato válido.",
        "No se permite prestar a un usuario inexistente ni un libro sin disponibilidad.",
        "Un préstamo devuelto no puede devolverse nuevamente.",
        "Préstamo, detalle y actualización de existencias se confirman o revierten como una unidad.",
    ]:
        add_bullet(doc, item)
    add_page_break(doc)

    add_heading(doc, "4. Casos de uso", 1)
    add_table(doc, 2, "Casos de uso principales", ["Caso", "Actor y flujo principal", "Postcondición"], [
        ("CU01 Administrar libro", "Encargado registra o edita código, ISBN, título, autor, categoría y cantidad; también puede buscar o eliminar.", "Catálogo y disponibilidad consistentes."),
        ("CU02 Administrar autor/categoría", "Encargado completa datos, guarda, edita o elimina cuando no existen relaciones.", "Catálogos auxiliares actualizados."),
        ("CU03 Administrar usuario", "Encargado registra documento, nombre, contacto y programa académico.", "Usuario disponible para préstamos."),
        ("CU04 Registrar préstamo", "Encargado elige usuario, libro disponible y fecha; el sistema ejecuta una transacción.", "Préstamo activo y disponibilidad reducida."),
        ("CU05 Registrar devolución", "Encargado selecciona un préstamo activo y confirma la devolución.", "Estado DEVUELTO y existencias restauradas."),
        ("CU06 Consultar", "Personal selecciona el tipo de consulta y presiona Consultar.", "Tabla actualizada sin modificar datos."),
    ], [1.25, 3.65, 1.75], 8)
    add_body(doc, "Los flujos alternos contemplan datos duplicados, referencias existentes, libros sin existencias, usuarios inexistentes, fechas inválidas y devoluciones repetidas. En todos estos casos la operación se rechaza sin dejar cambios parciales.")
    add_page_break(doc)

    add_heading(doc, "5. Diseño UML y arquitectura", 1)
    add_figure(doc, 1, "Arquitectura por capas del sistema", DIAGRAMS / "uml-arquitectura.png", "Diagrama UML de componentes. Las dependencias avanzan desde la presentación hasta SQL Server; Biblioteca.Entidades contiene los contratos compartidos.")
    add_page_break(doc)
    add_figure(doc, 2, "Modelo UML de clases del dominio", DIAGRAMS / "uml-clases.png", "Diagrama de clases con interfaz, clases abstractas, herencia y composición. La colección Detalles materializa la relación de composición del préstamo.")
    add_page_break(doc)
    add_figure(doc, 3, "Modelo lógico de la base de datos", DIAGRAMS / "uml-modelo-datos.png", "Relaciones y multiplicidades derivadas del script SQL. PK indica llave primaria; FK, llave foránea; UQ, restricción única.")
    add_page_break(doc)
    add_figure(doc, 4, "Secuencia para registrar un préstamo", DIAGRAMS / "uml-secuencia-prestamo.png", "El procedimiento almacenado valida disponibilidad y ejecuta INSERT y UPDATE dentro de una transacción. Ante un error, SQL Server ejecuta ROLLBACK.")
    add_page_break(doc)

    add_heading(doc, "6. Modelo de datos y diccionario", 1)
    add_body(doc, "La base BibliotecaDB contiene siete tablas. Aunque el enunciado exige como mínimo seis, se agregó Devoluciones para conservar evidencia independiente de la recepción. Las relaciones no usan borrado en cascada; así se evita eliminar información histórica accidentalmente.")
    add_table(doc, 3, "Diccionario de datos resumido", ["Tabla", "Llave y relaciones", "Campos principales", "Integridad"], [
        ("Autores", "PK AutorId", "Codigo, Nombre, Apellidos, Nacionalidad, FechaNacimiento", "Código único y fecha no futura"),
        ("Categorias", "PK CategoriaId", "Nombre, Descripcion, Activa", "Nombre único y obligatorio"),
        ("Usuarios", "PK UsuarioId", "Documento, Nombre, Apellidos, Telefono, Correo, ProgramaAcademico", "Documento/correo únicos y formatos CHECK"),
        ("Libros", "PK LibroId · FK AutorId · FK CategoriaId", "Codigo, ISBN, Titulo, CantidadTotal, CantidadDisponible", "Código/ISBN únicos y cantidades válidas"),
        ("Prestamos", "PK PrestamoId · FK UsuarioId", "Fechas y Estado", "Estado ACTIVO/DEVUELTO/VENCIDO"),
        ("DetallePrestamos", "PK DetallePrestamoId · FK PrestamoId · FK LibroId", "Cantidad", "Detalle único por préstamo y libro"),
        ("Devoluciones", "PK DevolucionId · FK/UQ PrestamoId", "FechaDevolucion, Observacion", "Una devolución por préstamo"),
    ], [1.15, 1.65, 2.55, 1.45], 7)
    add_heading(doc, "6.1 Vistas, índices y procedimientos", 2)
    add_body(doc, "Las vistas vw_LibrosDetalle y vw_HistorialPrestamos simplifican las consultas de interfaz. Los índices sobre título, autor, categoría y estado mejoran las búsquedas. sp_RegistrarPrestamo y sp_RegistrarDevolucion emplean XACT_ABORT, TRY/CATCH, bloqueos de actualización y transacciones explícitas (Microsoft, 2026a, 2026b).")
    add_page_break(doc)

    add_heading(doc, "7. Implementación de módulos", 1)
    modules = [
        ("7.1 Inicio", "Presenta títulos, ejemplares disponibles, usuarios y préstamos activos obtenidos de la base de datos."),
        ("7.2 Libros", "Implementa CRUD, búsqueda por código, título, autor o categoría y control de cantidades."),
        ("7.3 Autores", "Administra código, nombres, apellidos, nacionalidad y fecha de nacimiento."),
        ("7.4 Categorías", "Gestiona las categorías requeridas y permite crear categorías adicionales."),
        ("7.5 Usuarios", "Administra documento, nombres, teléfono, correo y programa académico."),
        ("7.6 Préstamos", "Muestra usuarios y libros disponibles; el registro se confirma mediante sp_RegistrarPrestamo."),
        ("7.7 Devoluciones", "Lista préstamos activos y restaura las existencias mediante sp_RegistrarDevolucion."),
        ("7.8 Consultas", "Ejecuta disponibles, prestados, usuarios con préstamos, historial, libros por categoría y total prestado."),
    ]
    for heading, text_value in modules:
        add_heading(doc, heading, 2)
        add_body(doc, text_value)

    add_heading(doc, "8. Validaciones y manejo de errores", 1)
    add_body(doc, "La solución aplica defensa en profundidad. Las entidades verifican obligatoriedad y coherencia; los servicios validan formatos y reglas cruzadas; los repositorios parametrizan cada consulta; SQL Server exige NOT NULL, UNIQUE, CHECK y FOREIGN KEY; y la interfaz traduce las excepciones conocidas a mensajes comprensibles.")
    for item in [
        "Dos préstamos simultáneos sobre el último ejemplar no pueden confirmarse porque el procedimiento bloquea la fila.",
        "Si falla el detalle, se revierten el encabezado y la reducción de disponibilidad.",
        "Una devolución repetida se rechaza porque el préstamo ya no está activo.",
        "La eliminación de un registro relacionado conserva el historial y produce un mensaje de integridad.",
        "Las entradas de búsqueda se transmiten como parámetros y no se concatenan como código SQL.",
    ]:
        add_bullet(doc, item)
    add_page_break(doc)

    add_heading(doc, "9. Evidencias reales del sistema", 1)
    add_body(doc, "Las figuras 5 a 12 son capturas auténticas de la aplicación Biblioteca.Presentacion ejecutada en Windows Forms y conectada a localhost\\SQLEXPRESS. No corresponden al modo de demostración. Los datos visibles provienen de BibliotecaDB: cuatro títulos, catorce ejemplares totales, doce disponibles, tres usuarios y dos préstamos activos.")
    capture_figures = [
        (5, "Tablero de inicio conectado a BibliotecaDB", required[0], "La aplicación muestra 4 títulos, 12 ejemplares disponibles, 3 usuarios y 2 préstamos activos."),
        (6, "Gestión real de libros", required[1], "La tabla presenta los cuatro títulos cargados, sus autores, categorías, cantidades y disponibilidad actual."),
        (7, "Gestión real de autores", required[2], "Se visualizan Robert Martin, Ian Sommerville y Andrew Tanenbaum con sus datos registrados."),
        (8, "Gestión real de categorías", required[3], "Se observan las siete categorías exigidas, incluida Otros con su descripción corregida."),
        (9, "Gestión real de usuarios", required[4], "La tabla contiene los tres usuarios de prueba almacenados en SQL Server."),
        (10, "Registro real de préstamos", required[5], "La base contiene dos préstamos activos asociados a LIB-001 y LIB-003."),
        (11, "Registro real de devoluciones", required[6], "El formulario lista los mismos préstamos activos y ofrece la operación transaccional de devolución."),
        (12, "Consulta real de libros por categoría", required[7], "La agregación confirma 2 títulos y 9 ejemplares en Programación, 1 y 2 en Bases de datos, y 1 y 3 en Redes."),
    ]
    for number, title, image_path, note in capture_figures:
        add_page_break(doc)
        add_figure(doc, number, title, image_path, note, 6.45)
    add_page_break(doc)

    add_heading(doc, "10. Pruebas de funcionamiento", 1)
    add_table(doc, 4, "Resultados de verificación", ["ID", "Escenario", "Resultado"], [
        ("T01", "Compilación de Entidades, Datos, Negocio y Presentación", "Aprobado en .NET Framework 4.8"),
        ("T02", "Ejecución conectada a localhost\\SQLEXPRESS", "Aprobado"),
        ("T03", "Carga de 4 libros, 3 autores, 7 categorías y 3 usuarios", "Aprobado"),
        ("T04", "Registro idempotente de 2 préstamos activos", "Aprobado"),
        ("T05", "Actualización de disponibilidad: 14 totales y 12 disponibles", "Aprobado"),
        ("T06", "Listado de préstamos en Devoluciones", "Aprobado"),
        ("T07", "Consulta de cantidad por categoría", "Aprobado"),
        ("T08", "Validaciones automáticas de dominio", "Aprobado"),
        ("T09", "Prueba transaccional de préstamo y devolución con ROLLBACK", "Incluida y reproducible"),
        ("T10", "Revisión visual de ocho capturas reales", "Aprobado"),
    ], [0.55, 4.45, 1.55], 8)
    add_body(doc, "La prueba sql/02_Pruebas_Integridad.sql crea un préstamo dentro de una transacción, comprueba la disminución de disponibilidad, registra la devolución, verifica su restauración y finaliza con ROLLBACK para no alterar los datos de demostración.")

    add_heading(doc, "11. Trazabilidad y control de versiones", 1)
    add_table(doc, 5, "Matriz de trazabilidad", ["Elemento", "Componentes", "Evidencia"], [
        ("CRUD", "Formularios, servicios y repositorios", "Figuras 6 a 9 y código fuente"),
        ("Préstamos", "PrestamosForm · PrestamoServicio · sp_RegistrarPrestamo", "Figura 10 y prueba SQL"),
        ("Devoluciones", "DevolucionesForm · sp_RegistrarDevolucion", "Figura 11 y prueba SQL"),
        ("Consultas", "ConsultasForm · ConsultaRepositorio · vistas", "Figura 12"),
        ("POO", "IValidable · EntidadBase · Persona · colección Detalles", "Figura 2"),
        ("Arquitectura", "Cuatro proyectos y SQL Server", "Figura 1"),
        ("Base relacional", "7 tablas, PK, FK, UQ, CHECK e índices", "Figura 3 y script SQL"),
        ("Git", "Repositorio local, README y commits", "docs/EVIDENCIA_GIT.md"),
    ], [1.25, 3.15, 2.2], 8)
    add_body(doc, "El repositorio local usa la rama main y conserva commits separados para dominio, datos y negocio, interfaz, base de datos, documentación y empaquetado. La publicación en GitHub debe realizarse desde la cuenta personal del estudiante, porque requiere sus credenciales.")
    add_page_break(doc)

    add_heading(doc, "12. Conclusiones y recomendaciones", 1)
    add_heading(doc, "12.1 Conclusiones", 2)
    conclusions = [
        "La aplicación centraliza la información y elimina la dependencia de registros manuales dispersos.",
        "La separación por capas reduce el acoplamiento y permite modificar la interfaz sin reescribir la persistencia.",
        "Las validaciones en varios niveles impiden que la integridad dependa de un único control visual.",
        "Las transacciones mantienen sincronizados préstamos, detalles, devoluciones y existencias.",
        "La herencia, la abstracción, el polimorfismo, el encapsulamiento y las colecciones evidencian la aplicación de POO.",
        "Las capturas reales demuestran la ejecución correcta de los módulos contra SQL Server.",
    ]
    for item in conclusions:
        add_bullet(doc, item)
    add_heading(doc, "12.2 Recomendaciones", 2)
    for item in [
        "Publicar el repositorio en GitHub y proteger la rama principal.",
        "Configurar copias de seguridad periódicas de BibliotecaDB.",
        "Usar una cadena de conexión externa por ambiente en una evolución del proyecto.",
        "Agregar autenticación, reservas, notificaciones y multas solo si el alcance futuro lo exige.",
        "Ejecutar las pruebas de integración antes de cada entrega o despliegue.",
    ]:
        add_bullet(doc, item)
    add_page_break(doc)

    add_heading(doc, "Referencias", 1)
    for ref in [
        "Ecma International. (2023). ECMA-334: C# language specification (7th ed.). https://ecma-international.org/publications-and-standards/standards/ecma-334/",
        "Git. (2026). Git documentation (Version 2.55.0). https://git-scm.com/docs/git",
        "Microsoft. (2025a, 7 de mayo). Windows Forms overview. Microsoft Learn. https://learn.microsoft.com/en-us/dotnet/desktop/winforms/overview/",
        "Microsoft. (2025b, 8 de agosto). SQL Server and ADO.NET. Microsoft Learn. https://learn.microsoft.com/en-us/dotnet/framework/data/adonet/sql/",
        "Microsoft. (2026a, 20 de julio). Transactions (Transact-SQL). Microsoft Learn. https://learn.microsoft.com/en-us/sql/t-sql/language-elements/transactions-transact-sql",
        "Microsoft. (2026b). Transaction locking and row versioning guide. Microsoft Learn. https://learn.microsoft.com/en-us/sql/relational-databases/sql-server-transaction-locking-and-row-versioning-guide",
    ]:
        add_reference(doc, ref)
    add_page_break(doc)

    add_heading(doc, "Apéndice A. Ejecución en Visual Studio", 1)
    steps = [
        "Instalar Visual Studio Community 2022 y seleccionar la carga de trabajo Desarrollo de escritorio con .NET.",
        "Instalar SQL Server Express y SQL Server Management Studio.",
        "En SSMS, conectarse a localhost\\SQLEXPRESS con Autenticación de Windows y ejecutar sql/01_Crear_BibliotecaDB.sql.",
        "Abrir SistemaGestionBiblioteca.sln en Visual Studio.",
        "Confirmar que Biblioteca.Presentacion sea el proyecto de inicio.",
        "Compilar la solución con Ctrl+Mayús+B y ejecutar con F5.",
        "Si cambia el nombre de la instancia, ajustar BibliotecaDb en src/Biblioteca.Presentacion/App.config.",
    ]
    for i, step in enumerate(steps, 1):
        add_body(doc, f"{i}. {step}", False, WD_ALIGN_PARAGRAPH.LEFT)
    add_heading(doc, "Apéndice B. Lista de entregables", 1)
    for item in [
        "Documento Markdown: docs/INFORME.md.",
        "Informe Word con APA 7: docs/INFORME_APA7.docx.",
        "Código fuente: SistemaGestionBiblioteca.sln y carpeta src.",
        "Script SQL: sql/01_Crear_BibliotecaDB.sql.",
        "Pruebas SQL: sql/02_Pruebas_Integridad.sql.",
        "Repositorio Git local, README.md y docs/EVIDENCIA_GIT.md.",
        "Ocho capturas reales en docs/capturas y cuatro diagramas UML en docs/diagramas.",
        "Archivo final SistemaGestionBiblioteca-Entrega.zip.",
    ]:
        add_bullet(doc, item)

    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal" and not paragraph.text:
            paragraph.paragraph_format.first_line_indent = Inches(0)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_report()
    print(output)
